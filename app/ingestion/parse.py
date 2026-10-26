"""Text extraction for PDF and DOCX.

PDF: text per page with pypdf. Pages with no text layer are skipped and reported
(OCR is not in this version). DOCX: paragraphs in document order. When the file
has tracked changes, two views are produced: 'accepted' (insertions kept,
deletions dropped) and 'original' (before the changes). Each chunk is labeled,
so deleted language is never presented as current terms.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader

from app.ingestion.validate import RejectedDocument, check_docx_archive

SCANNED_MIN_CHARS = 25


@dataclass(frozen=True)
class Block:
    text: str
    page: int | None
    is_heading: bool = False


@dataclass
class ParsedDocument:
    doc_type: str
    page_count: int | None
    views: dict[str, list[Block]]
    warnings: list[str] = field(default_factory=list)


def parse_pdf(path: Path, max_pages: int) -> ParsedDocument:
    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            raise RejectedDocument("encrypted", "Password-protected PDFs are not supported.")
        n = len(reader.pages)
    except RejectedDocument:
        raise
    except Exception as exc:
        raise RejectedDocument("corrupt_pdf", "The PDF could not be read.") from exc
    if n == 0:
        raise RejectedDocument("empty", "The PDF has no pages.")
    if n > max_pages:
        raise RejectedDocument("too_many_pages", f"The PDF has {n} pages; the limit is {max_pages}.")

    blocks: list[Block] = []
    scanned: list[int] = []
    for number, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        if len(text.strip()) < SCANNED_MIN_CHARS:
            scanned.append(number)
            continue
        blocks.extend(Block(line.strip(), number) for line in text.splitlines() if line.strip())

    if not blocks:
        raise RejectedDocument(
            "scanned_unsupported",
            "This PDF has no text layer. Scanned PDFs need OCR, which this version does not support.",
        )
    warnings = []
    if scanned:
        warnings.append(f"Pages without a text layer were skipped: {scanned}")
    return ParsedDocument("pdf", n, {"accepted": blocks}, warnings)


W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _has_ancestor(node, tag: str) -> bool:
    cur = node.getparent()
    while cur is not None:
        if cur.tag == tag:
            return True
        cur = cur.getparent()
    return False


def _paragraph_text(p, view: str) -> str:
    parts: list[str] = []
    for node in p.iter():
        if node.tag == qn("w:t"):
            if view == "accepted" and _has_ancestor(node, qn("w:del")):
                continue
            if view == "original" and _has_ancestor(node, qn("w:ins")):
                continue
            parts.append(node.text or "")
        elif node.tag == qn("w:delText") and view == "original":
            parts.append(node.text or "")
        elif node.tag == qn("w:tab") and node.getparent() is not None and node.getparent().tag == qn("w:r"):
            parts.append("\t")
    return "".join(parts).strip()


def _is_heading(p) -> bool:
    ppr = p.find(qn("w:pPr"))
    if ppr is None:
        return False
    style = ppr.find(qn("w:pStyle"))
    if style is None:
        return False
    return (style.get(qn("w:val")) or "").lower().startswith(("heading", "title"))


def parse_docx(path: Path, max_paragraphs: int = 20000) -> ParsedDocument:
    check_docx_archive(path)
    try:
        document = Document(str(path))
    except Exception as exc:
        raise RejectedDocument("corrupt_docx", "The DOCX file could not be read.") from exc

    body = document.element.body
    paragraphs = list(body.iter(qn("w:p")))
    if len(paragraphs) > max_paragraphs:
        raise RejectedDocument("too_large", "The document has too many paragraphs.")
    tracked = any(True for _ in body.iter(qn("w:ins"), qn("w:del")))

    views: dict[str, list[Block]] = {"accepted": []}
    if tracked:
        views["original"] = []
    for p in paragraphs:
        heading = _is_heading(p)
        for label in views:
            text = _paragraph_text(p, label)
            if text:
                views[label].append(Block(text, None, heading))

    if not views["accepted"]:
        raise RejectedDocument("empty", "The document has no text.")
    warnings = ["Document has tracked changes; both accepted and original views are indexed."] if tracked else []
    return ParsedDocument("docx", None, views, warnings)
