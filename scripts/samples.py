"""Build sample files for manual testing and for the test suite.

    python -m scripts.samples data/samples

Writes a DOCX with clauses and tracked changes, and a PDF with a text layer
plus a blank page to exercise the scanned-page path. The text is invented
for testing and is not drawn from any real agreement."""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

_next_revision_id = [1]


def _revision(kind: str, text: str):
    wrapper = OxmlElement("w:ins" if kind == "ins" else "w:del")
    wrapper.set(qn("w:id"), str(_next_revision_id[0]))
    _next_revision_id[0] += 1
    wrapper.set(qn("w:author"), "Reviewer")
    wrapper.set(qn("w:date"), "2026-01-15T09:00:00Z")
    run = OxmlElement("w:r")
    node = OxmlElement("w:delText" if kind == "del" else "w:t")
    node.set(qn("xml:space"), "preserve")
    node.text = text
    run.append(node)
    wrapper.append(run)
    return wrapper


def add_tracked_paragraph(doc, parts: list[tuple[str, str]]):
    paragraph = doc.add_paragraph()
    for kind, text in parts:
        if kind == "n":
            paragraph.add_run(text)
        else:
            paragraph._p.append(_revision(kind, text))
    return paragraph


def build_sample_docx(path: Path) -> Path:
    doc = Document()
    doc.add_heading("Master Services Agreement", level=1)
    doc.add_paragraph("ARTICLE 1 DEFINITIONS")
    doc.add_paragraph('1.1 "Confidential Information" means all non-public information disclosed by either party '
                      "in connection with this Agreement, whether marked or not.")
    add_tracked_paragraph(doc, [
        ("n", "1.2 Term. This Agreement starts on the Effective Date and continues for "),
        ("del", "two (2) years"),
        ("ins", "three (3) years"),
        ("n", " unless ended earlier under Section 9.3."),
    ])
    doc.add_paragraph("2 Limitation of Liability")
    doc.add_paragraph("2.1 Neither party is liable for indirect or consequential loss.")
    add_tracked_paragraph(doc, [
        ("n", "2.2 Cap. Each party's total liability is limited to "),
        ("del", "the fees paid in the prior six (6) months"),
        ("ins", "the fees paid in the prior twelve (12) months"),
        ("n", "."),
    ])
    doc.add_paragraph("9.3 Termination for Convenience. Either party may end this Agreement on ninety (90) days written notice.")
    doc.add_paragraph("(a) the notice must be given in writing;")
    doc.add_paragraph("(b) the notice takes effect on the date stated in it.")
    doc.save(path)
    return path


def build_sample_pdf(path: Path) -> Path:
    from reportlab.pdfgen import canvas

    canvas_ = canvas.Canvas(str(path))
    page1 = [
        "NON-DISCLOSURE AGREEMENT",
        "1 Definitions",
        "1.1 Confidential Information means information marked confidential by the disclosing party.",
        "2 Permitted Use",
        "2.1 The receiving party shall use Confidential Information only to evaluate the proposed transaction.",
    ]
    y = 780
    for line in page1:
        canvas_.drawString(72, y, line)
        y -= 18
    canvas_.showPage()
    canvas_.showPage()  # blank page: no text layer
    canvas_.save()
    return path


if __name__ == "__main__":
    out = Path(sys.argv[1] if len(sys.argv) > 1 else "data/samples")
    out.mkdir(parents=True, exist_ok=True)
    print(build_sample_docx(out / "sample_msa.docx"))
    print(build_sample_pdf(out / "sample_nda.pdf"))
